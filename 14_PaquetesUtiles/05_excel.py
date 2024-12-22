# =========================================================================
# 05_excel.py - Script para manejo de archivos Excel con openpyxl
# =========================================================================
"""
Este script demuestra cómo usar la librería `openpyxl` para cargar un archivo de Excel
utilizando diferentes opciones de ruta. Además, detalla los pasos para instalar y usar
el entorno virtual con `pipenv` antes de ejecutar el script.

Requisitos:
-----------
1. Tener instalado Python 3.x.
2. Instalar pipenv:
   >>> pip install pipenv
   
3. Instalar la librería openpyxl dentro de un entorno virtual:
   >>> pipenv install openpyxl

4. Crear un archivo de Excel (en este ejemplo: "ejemplo.xlsx") 
   y colocarlo dentro de la misma carpeta donde se ejecutará el script,
   o en otra ubicación cuya ruta conozcas.

5. Activar el entorno virtual antes de ejecutar el script:
   >>> pipenv shell

6. Posicionarte en la carpeta donde está el script (y el archivo Excel si vas a usar ruta relativa):
   >>> cd 14_PaquetesUtiles
"""

# =========================================================================
# SECCIÓN 1: IMPORTACIÓN Y CONFIGURACIÓN INICIAL
# =========================================================================

import openpyxl  # Importamos la biblioteca para manipulación de archivos Excel

# =========================================================================
# SECCIÓN 2: DIFERENTES FORMAS DE CARGAR UN ARCHIVO EXCEL
# =========================================================================

# Opción 1: Usar ruta relativa cuando el archivo "ejemplo.xlsx"
# está en la misma carpeta que 05_excel.py
wb = openpyxl.load_workbook("ejemplo.xlsx")

# Otras opciones disponibles (comentadas):
# -------------------------------------
# Opción 2: Usar ruta relativa indicando la carpeta
# wb = openpyxl.load_workbook("14_PaquetesUtiles/ejemplo.xlsx")

# Opción 3: Usar ruta absoluta (funciona desde cualquier ubicación)
# wb = openpyxl.load_workbook(r"C:\PythonCourse\14_PaquetesUtiles\ejemplo.xlsx")

# Opción 4: Determinar la ruta del script en tiempo de ejecución
# import os
# ruta_script = os.path.dirname(os.path.abspath(__file__))  # Obtiene la ruta del script actual
# ruta_excel = os.path.join(ruta_script, "ejemplo.xlsx")    # Combina las rutas de forma segura
# wb = openpyxl.load_workbook(ruta_excel)

# =========================================================================
# SECCIÓN 3: MANEJO DE HOJAS DE CÁLCULO
# =========================================================================

# Ver todas las hojas disponibles en el workbook
print("Hojas en este workbook:", wb.sheetnames)

# Ver información de una hoja específica
print(wb["Sheet1"])

# Diferentes formas de acceder a las hojas
hoja = wb.active      # Método 1: Obtener la hoja activa
correos = wb["Sheet1"]  # Método 2: Obtener hoja por nombre

# Crear una nueva hoja y cambiar su nombre
wb.create_sheet("NuevaHoja")
hoja3 = wb["NuevaHoja"]
hoja3.title = "NuevoNombre"

# Ver dimensiones de la hoja (filas y columnas)
print(
    hoja.max_row,     # Número total de filas
    hoja.max_column,  # Número total de columnas
)

# =========================================================================
# SECCIÓN 4: MANEJO DE CELDAS
# =========================================================================

# Método 1: Acceder a celda usando notación Excel
celda = hoja["A1"]
print(celda)          # Información de la celda
print(celda.value)    # Valor contenido en la celda

# Método 2: Acceder a celda de una hoja específica
print(correos["A1"].value)

# Modificar el valor de una celda
celda_correos = correos["A1"]
celda_correos.value = "Nombre Completo"
print(celda_correos.value)

# Esta es la salida esperada hasta este punto:
"""
PS C:\PythonCourse\14_PaquetesUtiles> & C:/Users/econo/.virtualenvs/14_PaquetesUtiles-vVT4Ex2a/Scripts/python.exe c:/PythonCourse/14_PaquetesUtiles/05_excel.py       
c:\PythonCourse\14_PaquetesUtiles\05_excel.py:1: SyntaxWarning: invalid escape sequence '\P'

Hojas en este workbook: ['Sheet1', 'Hoja1']
<Worksheet "Sheet1">
1 1
<Cell 'Hoja1'.A1>
c
Nombre
Nombre Completo
"""

# Método 3: Acceder a celda usando índices (basados en 1)
celda2 = correos.cell(row=2, column=1)
print(celda2.value)

# Obtener información completa de una celda
print(celda2.value,      # Contenido
      celda2.row,        # Número de fila
      celda2.column,     # Número de columna
      celda2.coordinate  # Coordenada estilo Excel (ej: A2)
      )

# =========================================================================
# SECCIÓN 5: EJEMPLOS DE ITERACIÓN SOBRE CELDAS
# =========================================================================

# Ejemplo 1: Iterar mostrando coordenadas (comentado)
"""
for fila in range(1, correos.max_row + 1):
    for columna in range(1, correos.max_column + 1):
        celda = correos.cell(row=fila, column=columna)
        print(fila, columna, celda.value)
    print()

Salida esperada:
1 1 Nombre Completo
1 2 Edad
1 3 Correo

2 1 Hugo
"""

# Ejemplo 2: Iterar mostrando datos en formato tabla (comentado)
"""
for fila in range(1, correos.max_row + 1):
    for columna in range(1, correos.max_column + 1):
        celda = correos.cell(row=fila, column=columna)
        print(celda.value, end=" ")
    print()

Salida esperada:
Nombre Completo Edad Correo
Hugo 32 hugo@hotmail.com
Paco  22 luis@hotmail.com
Luis 34 luis@hotmail.com
"""

# =========================================================================
# SECCIÓN 6: ACCESO A COLUMNAS Y FILAS COMPLETAS
# =========================================================================

# Ejemplo de acceso a una columna completa (comentado)
"""
columna = correos["A"]
print(columna)
# (<Cell 'Sheet1'.A1>, <Cell 'Sheet1'.A2>, <Cell 'Sheet1'.A3>, <Cell 'Sheet1'.A4>)

# Inspeccionar valores de la columna
for celda in columna:
    print(celda.value)

Salida esperada:
Nombre Completo
Hugo
Paco
Luis
"""

# Ejemplo de acceso a una fila completa (comentado)
"""
fila = correos[1]
print(fila)
# (<Cell 'Sheet1'.A1>, <Cell 'Sheet1'.B1>, <Cell 'Sheet1'.C1>)

for celda in fila:
    print(celda.value)

Salida esperada:
Nombre Completo 
Edad 
Correo
"""

# =========================================================================
# SECCIÓN 7: MODIFICACIONES A LA HOJA
# =========================================================================

# Agregar una nueva fila de datos
correos.append(["Ana", 23, "ana@hotmail.com"])
print(correos.rows)  # Muestra el generador de filas

# Guardar los cambios en el archivo original
wb.save("ejemplo.xlsx")

# Operaciones de estructura
# ------------------------
# Eliminar la primera fila
correos.delete_rows(1)

# Eliminar la primera columna
correos.delete_cols(1)

# Guardar los cambios en un nuevo archivo
wb.save("ejemplo_modificado.xlsx")
